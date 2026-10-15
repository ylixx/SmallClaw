#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
viewer_tool.py — 缩略图导出 + viewer.html 生成（二合一）

一条命令完成「交互调版式」链路：
  ① 可选：把渲染产物 --deck 导出 deck 缩略图（slide001.png…）
  ② 可选：把版式模板 --template 导出模板缩略图（tpl_001.png…，
    目录已有内容则跳过，--force-tpl 强制重新生成）
  ③ 生成 viewer.html（storyboard + patterns + layout_index + 缩略图）

用法：
  python viewer_tool.py storyboard.json -o viewer.html
  python viewer_tool.py storyboard.json --deck out.pptx --template template.pptx -o viewer.html

依赖：
  缩略图导出按 win32com(PowerPoint) → LibreOffice+PyMuPDF → Pillow 三级降级；
  无 Office 时 Pillow 兜底为白底占位。viewer.html 纯静态，无 CDN。
"""

import argparse
import json
import os
import sys
from html import escape

from PIL import Image

from checks import page_compat

FALLBACK_B64 = ("PHN2ZyB3aWR0aD0iNDAiIGhlaWdodD0iNDAiIHZpZXdCb3g9IjAgMCA0MCA0MCIg"
                "ZmlsbD0ibm9uZSIgeG1sbnM9Imh0dHA6Ly93d3cudzMub3JnLzIwMDAvc3ZnIj48"
                "cmVjdCB3aWR0aD0iNDAiIGhlaWdodD0iNDAiIGZpbGw9IiNGM0Y0RjYiLz48L3N2Zz4=")


def warn(msg):
    print(f"[warn] {msg}", file=sys.stderr)


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        raw = f.read()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        # 网页大模型常把 JSON 包在 ```json 代码块或前后缀文字里，自动提取
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start:end + 1])
        raise


def js_str(obj):
    """JSON 内联进 <script> 的安全序列化：转义 </script> 与 U+2028/2029。"""
    s = json.dumps(obj, ensure_ascii=False)
    return (s.replace("</", "<\\/")
             .replace("\u2028", "\\u2028")
             .replace("\u2029", "\\u2029"))


# ---------------------------------------------------------------- 缩略图导出
def _try_win32com(pptx_path, out_dir, width, height, prefix):
    """优先尝试 win32com.client 调 PowerPoint COM。"""
    try:
        import win32com.client
    except ImportError:
        return False
    try:
        ppt = win32com.client.Dispatch("PowerPoint.Application")
        # 新版 Office 禁止隐藏应用窗口（Visible=0 会抛错），改为可见或跳过
        try:
            ppt.Visible = 1
        except Exception:
            pass
        deck = ppt.Presentations.Open(os.path.abspath(pptx_path))
        n = deck.Slides.Count
        for i in range(1, n + 1):
            slide = deck.Slides(i)
            out_path = os.path.abspath(os.path.join(out_dir, f"{prefix}{i:03d}.png"))
            slide.Export(out_path, "PNG", width, height)
        deck.Close()
        ppt.Quit()
        return True
    except Exception as e:
        warn(f"win32com COM 导出失败: {e}")
        try:
            deck.Close()
            ppt.Quit()
        except Exception:
            pass
        return False


def _try_libreoffice(pptx_path, out_dir, width, height, prefix):
    """备选：LibreOffice 转为 PDF + PyMuPDF 拆页。"""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return False
    try:
        soffice = "soffice"  # 默认命令名
        if sys.platform == "win32":
            soffice = "soffice.exe"
        # 转为 PDF
        pdf_path = os.path.join(out_dir, "temp.pdf")
        cmd = f'"{soffice}" --headless --convert-to pdf --outdir "{out_dir}" "{os.path.abspath(pptx_path)}"'
        os.system(cmd)
        if not os.path.exists(pdf_path):
            raise FileNotFoundError("LibreOffice 未生成 PDF")
        # 拆页为 PNG
        doc = fitz.open(pdf_path)
        n = doc.page_count
        for i in range(n):
            page = doc.load_page(i)
            pix = page.get_pixmap(matrix=fitz.Matrix(width/72, height/72))
            out_path = os.path.join(out_dir, f"{prefix}{i+1:03d}.png")
            pix.save(out_path)
        doc.close()
        os.remove(pdf_path)
        return True
    except Exception as e:
        warn(f"LibreOffice+PyMuPDF 导出失败: {e}")
        return False


def _try_pillow(pptx_path, out_dir, width, height, prefix):
    """纯 Pillow 离线兜底：python-pptx 读取每页，PIL 渲染为 PNG。"""
    try:
        from pptx import Presentation
    except ImportError:
        warn("Pillow 离线兜底需要 python-pptx")
        return False
    try:
        prs = Presentation(pptx_path)
        os.makedirs(out_dir, exist_ok=True)
        for i, slide in enumerate(prs.slides, 1):
            # 简单渲染：取第一张图片（如果有），否则白底黑字占位
            img = None
            for shape in slide.shapes:
                if hasattr(shape, 'image') and shape.image:
                    img = shape.image
                    break
            if img and hasattr(img, 'blob'):
                from io import BytesIO
                img_bytes = BytesIO(img.blob)
                img_pil = Image.open(img_bytes)
                img_pil = img_pil.resize((width, height), Image.LANCZOS)
            else:
                img_pil = Image.new("RGB", (width, height), "white")
                from PIL import ImageDraw, ImageFont
                draw = ImageDraw.Draw(img_pil)
                try:
                    font = ImageFont.truetype("arial.ttf", 24)
                except Exception:
                    font = ImageFont.load_default()
                draw.text((10, height - 30), f"Page {i}", fill="black", font=font)
            out_path = os.path.join(out_dir, f"{prefix}{i:03d}.png")
            img_pil.save(out_path, "PNG")
        return True
    except Exception as e:
        warn(f"Pillow 离线兜底失败: {e}")
        return False


def _count_slides(pptx_path):
    """读取源 PPTX 页数，用于导出后对比；失败返回 None 不告警。"""
    try:
        from pptx import Presentation
        return len(Presentation(pptx_path).slides)
    except Exception:
        return None


def export_thumbs(pptx_path, out_dir, width=1280, height=720, prefix="slide",
                  pattern_only=False):
    """把一个 .pptx 的每页导出为 PNG。返回导出的页数。"""
    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"PPTX 文件不存在: {pptx_path}")
    os.makedirs(out_dir, exist_ok=True)
    if pattern_only:
        prefix = "tpl_"
    # 清掉同前缀旧缩略图（运行时生成物），避免 viewer 混入过期页面
    for f in os.listdir(out_dir):
        if f.startswith(prefix) and f.endswith(".png"):
            os.remove(os.path.join(out_dir, f))

    success = False
    if _try_win32com(pptx_path, out_dir, width, height, prefix):
        success = True
    elif _try_libreoffice(pptx_path, out_dir, width, height, prefix):
        success = True
    elif _try_pillow(pptx_path, out_dir, width, height, prefix):
        success = True
    else:
        raise RuntimeError("所有导出方案均不可用。请安装 PowerPoint (win32com)、"
                           "LibreOffice+PyMuPDF 或确保 python-pptx+Pillow 可用。")

    files = [f for f in os.listdir(out_dir) if f.startswith(prefix) and f.endswith(".png")]
    expected = _count_slides(pptx_path)
    if expected is not None and len(files) != expected:
        warn(f"导出页数异常：预期 {expected} 页，实际 {len(files)} 页")
    return len(files)


# ---------------------------------------------------------------- viewer.html
VIEWER_HTML = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title} - 版式切换器</title>
<style>
body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; margin: 0; background: #f5f5f5; }}
.toolbar {{ position: sticky; top: 0; background: white; padding: 10px 20px; box-shadow: 0 2px 4px rgba(0,0,0,.1); z-index: 100; display: flex; align-items: center; gap: 16px; flex-wrap: wrap; }}
.toolbar h1 {{ margin: 0; font-size: 18px; }}
.toolbar .stats {{ color: #666; font-size: 13px; }}
.toolbar .bad {{ color: #d32f2f; font-weight: 600; }}
.toolbar .ok {{ color: #2e7d32; font-weight: 600; }}
.toolbar button {{ background: #007bff; color: white; border: none; padding: 6px 14px; border-radius: 4px; cursor: pointer; }}
.toolbar button:hover {{ background: #0056b3; }}
.grid {{ display: grid; grid-template-columns: repeat(auto-fill, minmax(200px, 1fr)); gap: 16px; padding: 20px; }}
.card {{ background: white; border-radius: 8px; overflow: hidden; box-shadow: 0 2px 8px rgba(0,0,0,.1); cursor: pointer; transition: transform .15s, box-shadow .15s; user-select: none; }}
.card:hover {{ transform: translateY(-2px); box-shadow: 0 4px 14px rgba(0,0,0,.16); }}
.thumb-wrap {{ position: relative; }}
.card img {{ width: 100%; height: 120px; object-fit: cover; display: block; background: #eee; }}
.dup-badge {{ position: absolute; top: 6px; right: 6px; background: #d32f2f; color: white; font-size: 12px; font-weight: 700; padding: 2px 7px; border-radius: 10px; }}
.card-info {{ padding: 10px 12px; }}
.card-page {{ font-size: 12px; color: #666; margin-bottom: 4px; }}
.card-pattern {{ font-weight: bold; font-size: 14px; word-break: break-all; }}
.context-menu {{ position: fixed; background: white; border: 1px solid #ddd; border-radius: 6px; box-shadow: 0 6px 20px rgba(0,0,0,.18); padding: 4px 0; z-index: 1000; max-height: 82vh; overflow-y: auto; display: none; min-width: 300px; }}
.menu-top {{ position: sticky; top: 0; background: white; padding: 6px 12px; font-size: 12px; color: #888; border-bottom: 1px solid #eee; display: flex; align-items: center; justify-content: space-between; }}
.menu-back {{ cursor: pointer; color: #007bff; font-weight: 600; }}
.menu-group-title {{ padding: 6px 12px; font-size: 12px; color: #888; font-weight: 700; background: #fafafa; border-bottom: 1px solid #eee; }}
.menu-item {{ padding: 8px 12px; cursor: pointer; display: flex; align-items: center; }}
.menu-item:hover {{ background: #f0f4ff; }}
.menu-item.current {{ background: #e3f2fd; }}
.menu-item.incompatible {{ opacity: .62; }}
.menu-item img {{ width: 44px; height: 44px; object-fit: cover; margin-right: 12px; border-radius: 4px; background: #f3f4f6; flex: none; }}
.menu-item-content {{ flex: 1; min-width: 0; }}
.menu-item-name {{ font-weight: bold; font-size: 13px; }}
.menu-item-desc {{ font-size: 12px; color: #666; margin-top: 1px; }}
.menu-item-compat {{ font-size: 11px; margin-top: 3px; }}
.compat-ok {{ color: #2e7d32; }}
.compat-warn {{ color: #b26a00; }}
.compat-bad {{ color: #d32f2f; }}
.menu-separator {{ border-top: 1px solid #eee; margin: 4px 0; }}
.menu-expander {{ padding: 8px 12px; font-style: italic; color: #666; cursor: pointer; }}
.menu-expander:hover {{ background: #f0f4ff; }}
.hidden {{ display: none; }}
.toast {{ position: fixed; left: 50%; bottom: 32px; transform: translateX(-50%); background: rgba(0,0,0,.82); color: white; padding: 10px 20px; border-radius: 20px; font-size: 14px; z-index: 2000; opacity: 0; transition: opacity .25s; pointer-events: none; }}
.modal-mask {{ position: fixed; inset: 0; background: rgba(0,0,0,.4); z-index: 1500; display: none; align-items: center; justify-content: center; }}
.modal {{ background: white; border-radius: 8px; max-width: 640px; width: 90%; max-height: 78vh; overflow-y: auto; padding: 20px; }}
.modal h3 {{ margin-top: 0; }}
.modal ul {{ margin: 8px 0 0; padding-left: 20px; }}
.modal .warn-item {{ color: #b26a00; }}
.modal .bad-item {{ color: #d32f2f; }}
.modal-actions {{ margin-top: 16px; text-align: right; }}
.modal button {{ margin-left: 8px; padding: 6px 14px; border-radius: 4px; cursor: pointer; border: 1px solid #ccc; background: white; }}
.modal button.primary {{ background: #007bff; color: white; border-color: #007bff; }}
</style>
</head>
<body>
<div class="toolbar">
    <h1>{title}</h1>
    <div class="stats" id="stats"></div>
    <button onclick="exportStoryboard()">导出 storyboard.json</button>
</div>
<div class="grid" id="grid">
{cards_html}
</div>
<div class="context-menu" id="contextMenu"></div>
<div class="toast" id="toast"></div>
<div class="modal-mask" id="modalMask" onclick="if(event.target===this)closeModal()">
    <div class="modal" id="modalBody"></div>
</div>
<script>
const FALLBACK = 'data:image/svg+xml;base64,{FALLBACK_B64}';
const storyboard = {JS_STORYBOARD};
const intentGroups = {JS_GROUPS};
const PAGE_COMPAT = {JS_PAGECOMPAT};
const ESC = (s) => String(s).replace(/[&<>"']/g, c => ({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}}[c]));
let currentPage = null;
let menuMode = 'intent';

function imgFail(img) {{
    if (img.dataset.failed) return;
    img.dataset.failed = '1';
    img.style.opacity = .3;
    img.src = FALLBACK;
}}

function findIntentForPattern(pattern) {{
    for (const intent in intentGroups) {{
        if (intentGroups[intent].items.some(it => it.name === pattern)) return intent;
    }}
    return null;
}}

/* 兼容性数据由服务端预计算注入（规则唯一来源：checks.py page_compat） */
function compatFor(pageIdx, patternName) {{
    const m = PAGE_COMPAT[pageIdx];
    return (m && m[patternName]) || {{bads: [], warns: [], ok: true}};
}}

function renderGroup(intent, group, page) {{
    let h = '<div class="menu-group-title">' + ESC(intent) + ' · ' + ESC(group.desc || '') + '</div>';
    for (const item of group.items) {{
        const c = compatFor(currentPage, item.name);
        const cls = ['menu-item'];
        if (item.name === page.pattern) cls.push('current');
        else if (c.bads.length) cls.push('incompatible');
        const compat = c.ok ? '<span class="compat-ok">✅ 兼容</span>'
            : (c.bads.length ? '<span class="compat-bad">❌ ' + ESC(c.bads.join('; ')) + '</span>'
                             : '<span class="compat-warn">⚠ ' + ESC(c.warns.join('; ')) + '</span>');
        h += '<div class="' + cls.join(' ') + '" data-pat="' + ESC(item.name) + '" onclick="selectPattern(this.dataset.pat)">'
            + '<img src="' + (item.tpl_img || FALLBACK) + '" alt="' + ESC(item.name) + '" onerror="imgFail(this)">'
            + '<div class="menu-item-content">'
            + '<div class="menu-item-name">' + ESC(item.name) + '</div>'
            + '<div class="menu-item-desc">' + ESC(item.desc) + '</div>'
            + '<div class="menu-item-compat">' + compat + '</div>'
            + '</div></div>';
    }}
    return h;
}}

function openMenu(event, pageIdx) {{
    event.preventDefault();
    event.stopPropagation();
    currentPage = pageIdx;
    menuMode = 'intent';
    const page = storyboard.pages[pageIdx - 1];
    const intent = page._intent || findIntentForPattern(page.pattern);
    const menu = document.getElementById('contextMenu');
    let html = '<div class="menu-top">第 ' + pageIdx + ' 页 · 切换版式'
        + '<span class="menu-back" onclick="event.stopPropagation();closeMenu()">✕</span></div>';
    if (intent && intentGroups[intent]) {{
        html += renderGroup(intent, intentGroups[intent], page);
        html += '<div class="menu-separator"></div>';
    }}
    html += '<div class="menu-expander" onclick="showAllIntents(event)">📋 更多版式…</div>';
    menu.innerHTML = html;
    positionMenu(menu, event);
}}

function showAllIntents(event) {{
    event.stopPropagation();
    menuMode = 'all';
    const page = storyboard.pages[currentPage - 1];
    const menu = document.getElementById('contextMenu');
    let html = '<div class="menu-top"><span class="menu-back" onclick="event.stopPropagation();backToIntent()">← 返回</span>'
        + '<span class="menu-back" onclick="event.stopPropagation();closeMenu()">✕</span></div>';
    for (const intent in intentGroups) {{
        html += renderGroup(intent, intentGroups[intent], page);
        html += '<div class="menu-separator"></div>';
    }}
    menu.innerHTML = html;
    positionMenu(menu, event);
}}

function backToIntent() {{
    const page = storyboard.pages[currentPage - 1];
    const intent = page._intent || findIntentForPattern(page.pattern);
    const menu = document.getElementById('contextMenu');
    let html = '<div class="menu-top">第 ' + currentPage + ' 页 · 切换版式'
        + '<span class="menu-back" onclick="event.stopPropagation();closeMenu()">✕</span></div>';
    if (intent && intentGroups[intent]) {{
        html += renderGroup(intent, intentGroups[intent], page);
        html += '<div class="menu-separator"></div>';
    }}
    html += '<div class="menu-expander" onclick="showAllIntents(event)">📋 更多版式…</div>';
    menu.innerHTML = html;
    menuMode = 'intent';
    positionMenu(menu, {{clientX: menu.style.left, clientY: menu.style.top}});
}}

function positionMenu(menu, event) {{
    menu.style.display = 'block';
    const x = (typeof event.clientX === 'number') ? event.clientX : parseFloat(menu.style.left) || 0;
    const y = (typeof event.clientY === 'number') ? event.clientY : parseFloat(menu.style.top) || 0;
    const mw = menu.offsetWidth, mh = menu.offsetHeight;
    let left = x, top = y;
    if (left + mw > window.innerWidth - 8) left = Math.max(8, window.innerWidth - mw - 8);
    if (top + mh > window.innerHeight - 8) top = Math.max(8, window.innerHeight - mh - 8);
    menu.style.left = left + 'px';
    menu.style.top = top + 'px';
}}

function selectPattern(patternName) {{
    if (currentPage) {{
        const page = storyboard.pages[currentPage - 1];
        page.pattern = patternName;
        const card = document.querySelector('.card[data-page="' + currentPage + '"]');
        if (card) {{
            card.dataset.pattern = patternName;
            const nameEl = card.querySelector('.card-pattern');
            if (nameEl) nameEl.textContent = patternName;
            const img = card.querySelector('img');
            const tpl = findTplImg(patternName);
            if (img && tpl) {{ img.dataset.failed = ''; img.style.opacity = 1; img.src = tpl; }}
        }}
        refreshStats();
        toast('已切换为 ' + patternName);
    }}
    closeMenu();
}}

function findTplImg(patternName) {{
    for (const intent in intentGroups) {{
        for (const item of intentGroups[intent].items) {{
            if (item.name === patternName && item.tpl_img) return item.tpl_img;
        }}
    }}
    return null;
}}

function closeMenu() {{
    const menu = document.getElementById('contextMenu');
    if (menu) menu.style.display = 'none';
    currentPage = null;
}}

function toast(msg) {{
    const t = document.getElementById('toast');
    t.textContent = msg;
    t.style.opacity = 1;
    clearTimeout(t._timer);
    t._timer = setTimeout(() => {{ t.style.opacity = 0; }}, 1600);
}}

function refreshStats() {{
    const STRUCTURAL = {structural_js};
    const usage = {{}};
    let adjacent = 0;
    for (let i = 0; i < storyboard.pages.length; i++) {{
        const p = storyboard.pages[i].pattern;
        usage[p] = (usage[p] || 0) + 1;
        if (i > 0 && storyboard.pages[i - 1].pattern === p &&
            STRUCTURAL.indexOf(p) < 0) adjacent++;
    }}
    let dupPatterns = 0, dupPages = 0;
    for (const p in usage) {{
        if (usage[p] > 1 && STRUCTURAL.indexOf(p) < 0) {{
            dupPatterns++;
            dupPages += usage[p];
        }}
    }}
    const uniq = Object.keys(usage).length;
    const el = document.getElementById('stats');
    let txt = '共 ' + storyboard.pages.length + ' 页 · 版式 ' + uniq + ' 种';
    if (adjacent > 0) txt += ' <span class="bad">· 相邻重复 ' + adjacent + ' 处</span>';
    else txt += ' <span class="ok">· 无相邻重复</span>';
    if (dupPatterns > 0) txt += ' <span class="bad">· 重复使用 ' + dupPatterns + ' 种/占 ' + dupPages + ' 页</span>';
    el.innerHTML = txt;
    document.querySelectorAll('.card').forEach(card => {{
        const p = storyboard.pages[parseInt(card.dataset.page, 10) - 1].pattern;
        const n = usage[p] || 0;
        let badge = card.querySelector('.dup-badge');
        if (n > 1 && STRUCTURAL.indexOf(p) < 0) {{
            if (!badge) {{
                badge = document.createElement('span');
                badge.className = 'dup-badge';
                card.querySelector('.thumb-wrap').appendChild(badge);
            }}
            badge.textContent = '×' + n;
        }} else if (badge) {{
            badge.remove();
        }}
    }});
}}

function collectIssues() {{
    const issues = [];
    for (let i = 0; i < storyboard.pages.length; i++) {{
        const page = storyboard.pages[i];
        const item = findPatternItem(page.pattern);
        if (item) {{
            const c = compatFor(i + 1, page.pattern);
            if (c.bads.length || c.warns.length) {{
                issues.push('第 ' + (i + 1) + ' 页 [' + page.pattern + ']: '
                    + c.bads.concat(c.warns).join('; '));
            }}
        }} else {{
            issues.push('第 ' + (i + 1) + ' 页: 版式 "' + page.pattern + '" 不在版式库中');
        }}
    }}
    return issues;
}}

function findPatternItem(name) {{
    for (const intent in intentGroups) {{
        for (const item of intentGroups[intent].items) {{
            if (item.name === name) return item;
        }}
    }}
    return null;
}}

function exportStoryboard() {{
    const issues = collectIssues();
    const modal = document.getElementById('modalMask');
    const body = document.getElementById('modalBody');
    let h = '<h3>导出前体检</h3>';
    if (issues.length) {{
        h += '<ul>';
        for (const it of issues) h += '<li class="bad-item">' + ESC(it) + '</li>';
        h += '</ul><p>以上页面建议修正后再导出（渲染器会兜底替换未知版式，但不兼容槽位会留空）。</p>';
    }} else {{
        h += '<p style="color:#2e7d32">✅ 全部 ' + storyboard.pages.length + ' 页版式兼容，无缺失。</p>';
    }}
    h += '<div class="modal-actions"><button onclick="closeModal()">取消</button>'
        + '<button class="primary" onclick="doDownload()">仍要导出</button></div>';
    body.innerHTML = h;
    modal.style.display = 'flex';
}}

function closeModal() {{
    document.getElementById('modalMask').style.display = 'none';
}}

function doDownload() {{
    closeModal();
    const blob = new Blob([JSON.stringify(storyboard, null, 2)], {{type: 'application/json'}});
    const url = URL.createObjectURL(blob);
    const a = document.createElement('a');
    a.href = url;
    a.download = 'storyboard.json';
    a.click();
    URL.revokeObjectURL(url);
    toast('已导出 storyboard.json');
}}

document.addEventListener('click', (e) => {{
    if (!e.target.closest('.card') && !e.target.closest('.context-menu')) closeMenu();
}});
document.addEventListener('keydown', (e) => {{
    if (e.key === 'Escape') {{ closeMenu(); closeModal(); }}
}});
document.addEventListener('contextmenu', (e) => {{
    if (!e.target.closest('.card') && !e.target.closest('.context-menu')) e.preventDefault();
}});
refreshStats();
</script>
</body>
</html>
"""


def build_viewer(sb, li, patterns, thumbs, tpl_thumbs, out_path):
    """生成 viewer.html（逻辑同原 make_viewer.py）。"""
    li = {k: v for k, v in li.items() if not k.startswith("_")}
    # 结构页 id：优先读版式库 _roles 声明，缺省回退默认清单
    struct_ids = set()
    roles = patterns.get("_roles")
    if isinstance(roles, dict):
        for v in roles.values():
            if isinstance(v, list):
                struct_ids.update(x for x in v if isinstance(x, str))
    if not struct_ids:
        struct_ids = {"cover", "toc", "section", "thanks", "refs"}
    patterns = {k: v for k, v in patterns.items() if not k.startswith("_")}

    pattern_to_idx = {}
    for name, pat in patterns.items():
        if "slide" in pat:
            pattern_to_idx[name] = pat["slide"]

    pages = sb.get("pages", [])
    title = escape(sb.get("title", "Storyboard Viewer"))

    intent_groups = {}
    for intent, data in li.items():
        items = []
        for cand_name in data.get("candidates", []):
            cand_pat = patterns.get(cand_name)
            if not cand_pat:
                continue
            tpl_idx = pattern_to_idx.get(cand_name)
            items.append({
                "name": cand_name,
                "desc": cand_pat.get("desc", ""),
                "tpl_img": (f"{tpl_thumbs}/tpl_{tpl_idx:03d}.png"
                            if tpl_idx else ""),
            })
        intent_groups[intent] = {"desc": data.get("desc", ""), "items": items}

    # 服务端预计算：每页 × 每个候选版式的兼容结果（规则唯一来源 = checks.py）
    all_cands = sorted({it["name"] for g in intent_groups.values()
                        for it in g["items"]})
    page_compat_map = {}
    for i, p in enumerate(pages, 1):
        page_compat_map[str(i)] = {
            cname: page_compat(p, patterns.get(cname, {}))
            for cname in all_cands
        }

    usage = {}
    for p in pages:
        usage[p.get("pattern", "")] = usage.get(p.get("pattern", ""), 0) + 1
    adjacent = sum(1 for i in range(1, len(pages))
                   if pages[i].get("pattern") == pages[i - 1].get("pattern")
                   and pages[i].get("pattern") not in struct_ids)

    def card(i, page):
        pattern = page.get("pattern", "unknown")
        dup = usage.get(pattern, 0)
        badge = ""
        if dup > 1 and pattern not in struct_ids:
            badge = f'<span class="dup-badge">×{dup}</span>'
        return (
            f'<div class="card" data-page="{i}" data-pattern="{pattern}" '
            f'oncontextmenu="openMenu(event,{i});return false;" '
            f'onclick="openMenu(event,{i})">'
            f'<div class="thumb-wrap"><img src="{thumbs}/slide{i:03d}.png" '
            f'alt="第{i}页" onerror="imgFail(this)">{badge}</div>'
            f'<div class="card-info">'
            f'<div class="card-page">第 {i} 页</div>'
            f'<div class="card-pattern">{pattern}</div>'
            f'</div></div>'
        )

    cards_html = "\n".join(card(i, p) for i, p in enumerate(pages, 1))
    structural_js = json.dumps(sorted(struct_ids), ensure_ascii=False)

    html = VIEWER_HTML.format(
        title=title,
        cards_html=cards_html,
        structural_js=structural_js,
        JS_STORYBOARD=js_str(sb),
        JS_GROUPS=js_str(intent_groups),
        JS_PAGECOMPAT=js_str(page_compat_map),
        FALLBACK_B64=FALLBACK_B64,
    )

    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"✅ 已生成 {out_path}")


def _default_asset(name):
    """兼容 skill 布局(../assets/)与独立布局(./ 或 ./assets/)的资产路径探测。"""
    here = os.path.dirname(os.path.abspath(__file__))
    root = os.path.dirname(here)
    for cand in (os.path.join(root, "assets", name),
                 os.path.join(here, "assets", name),
                 os.path.join(here, name)):
        if os.path.exists(cand):
            return cand
    return os.path.join(root, "assets", name)


def main():
    ap = argparse.ArgumentParser(
        description="缩略图导出 + viewer.html 生成（二合一）")
    ap.add_argument("storyboard", help="storyboard.json 路径")
    ap.add_argument("-o", "--out", default="viewer.html",
                    help="输出 HTML 路径（默认 viewer.html）")
    ap.add_argument("-p", "--patterns", default=_default_asset("patterns.json"),
                    help="patterns.json 路径（默认自动探测 assets/ 或脚本目录）")
    ap.add_argument("-l", "--layout", default=_default_asset("layout_index.json"),
                    help="layout_index.json 路径（默认自动探测 assets/ 或脚本目录）")
    ap.add_argument("-t", "--thumbs", default="deck_thumbs",
                    help="deck 缩略图目录（默认 deck_thumbs/）")
    ap.add_argument("--tpl-thumbs", default="tpl_thumbs",
                    help="模板缩略图目录（默认 tpl_thumbs/）")
    ap.add_argument("--deck", help="渲染产物 pptx：自动导出 deck 缩略图")
    ap.add_argument("--template", help="版式模板 pptx：自动导出 tpl 缩略图")
    ap.add_argument("--force-tpl", action="store_true",
                    help="强制重新生成模板缩略图（默认目录有内容则跳过）")
    ap.add_argument("-w", "--width", type=int, default=1280,
                    help="PNG 宽度（默认 1280）")
    ap.add_argument("--height", type=int, default=720,
                    help="PNG 高度（默认 720）")
    args = ap.parse_args()

    # ① deck 缩略图
    if args.deck:
        n = export_thumbs(args.deck, args.thumbs, args.width, args.height, "slide")
        print(f"✅ 已导出 {n} 页 deck 缩略图到 {args.thumbs}/")

    # ② 模板缩略图（目录已有内容则跳过，--force-tpl 强制）
    if args.template:
        need = args.force_tpl or not (
            os.path.isdir(args.tpl_thumbs) and os.listdir(args.tpl_thumbs))
        if need:
            n = export_thumbs(args.template, args.tpl_thumbs,
                              args.width, args.height, "tpl", pattern_only=True)
            print(f"✅ 已导出 {n} 页模板缩略图到 {args.tpl_thumbs}/")
        else:
            print(f"· 模板缩略图已存在，跳过（--force-tpl 强制重新生成）")

    # ③ 生成 viewer.html
    sb = load_json(args.storyboard)
    li = load_json(args.layout)
    patterns = load_json(args.patterns)
    build_viewer(sb, li, patterns, args.thumbs, args.tpl_thumbs, args.out)


if __name__ == "__main__":
    main()
